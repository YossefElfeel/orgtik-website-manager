"""Generates the seed binaries for the local file-manager mock (local-dev/apps/files/assets).

Run once (already done; outputs are checked in):  python generate_assets.py
Needs Python 3 with Pillow, PyMuPDF (fitz) and openpyxl, plus Windows fonts (Segoe UI, Dubai).
Everything here is fictional demo content for Orgtik.
"""
import html
import io
import json
import math
import os
import random
import shutil
import sys
import zipfile
from pathlib import Path

import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import docs_content as C  # noqa: E402
import h264  # noqa: E402

ROOT = HERE.parent                       # local-dev/apps/files
ASSETS = ROOT / "assets"
PUBLIC = ROOT.parents[2] / "file-manager-dashboard-main" / "public"
FONTS = Path("C:/Windows/Fonts")
PURPLE = "#8135CD"
MANIFEST = {}


def font(name, size):
    return ImageFont.truetype(str(FONTS / name), size)


def reset_dir(p):
    if p.exists():
        shutil.rmtree(p)
    p.mkdir(parents=True)


def rel(p):
    return p.relative_to(ASSETS).as_posix()


def register(key, path, mime, preview=None):
    MANIFEST[key] = {"file": rel(path), "mimeType": mime, "size": path.stat().st_size, "preview": preview or {"kind": "unsupported"}}


# ============================================================== PDF (PyMuPDF Story)
CSS = """
@font-face {font-family: segoe; src: url(segoeui.ttf);}
@font-face {font-family: segoe; font-weight: bold; src: url(segoeuib.ttf);}
@font-face {font-family: dubai; src: url(DUBAI-REGULAR.TTF);}
@font-face {font-family: dubai; font-weight: bold; src: url(DUBAI-BOLD.TTF);}
body {font-family: %(font)s; font-size: 10pt; color: #1f2937; line-height: 1.35;}
.band {background-color: #8135cd; color: #ffffff; padding: 7pt 12pt; font-size: 9pt;}
.band b {letter-spacing: 2pt;}
h1 {color: #8135cd; font-size: 17pt; margin: 14pt 0 6pt 0;}
h2 {color: #1e1b4b; font-size: 12pt; margin: 12pt 0 4pt 0;}
p {margin: 0 0 7pt 0;}
ul {margin: 0 0 8pt 0;}
li {margin-bottom: 2pt;}
table {width: 100%%; border-collapse: collapse; margin: 4pt 0 10pt 0;}
th {background-color: #f3e8ff; color: #3b0764; text-align: left; padding: 4pt 5pt; border: 1px solid #d8b4fe; font-size: 9pt;}
td {padding: 4pt 5pt; border: 1px solid #e5e7eb; font-size: 9pt;}
tr.last td {font-weight: bold; background-color: #faf5ff;}
table.kv td {border: none; border-bottom: 1px solid #f1f5f9; padding: 3pt 5pt;}
table.kv td.k {color: #6b7280;}
.note {color: #6b7280; font-size: 9pt; background-color: #f9fafb; padding: 6pt 8pt;}
.ar {font-family: dubai; font-size: 11pt;}
h1.ar {font-family: dubai; color: #8135cd; font-size: 17pt;}
h2.ar {font-family: dubai; color: #1e1b4b; font-size: 13pt;}
.draft {color: #dc2626; font-size: 20pt; font-weight: bold; letter-spacing: 3pt; border: 2px solid #dc2626; padding: 4pt 8pt; text-align: center;}
table.sign td {border: none; padding: 26pt 5pt 2pt 5pt;}
.line {border-top: 1px solid #9ca3af; padding-top: 3pt; font-size: 9pt; color: #374151;}
.swatch {color: #ffffff; font-weight: bold; padding: 14pt 6pt;}
"""

COVER_CSS = """
.cover {background-color: #8135cd; color: #ffffff; padding: 150pt 48pt 170pt 48pt;}
.kicker {font-size: 11pt; letter-spacing: 4pt; font-weight: bold; margin-bottom: 20pt;}
.ctitle {font-size: 34pt; font-weight: bold; line-height: 1.1; margin-bottom: 12pt;}
.csub {font-size: 13pt;}
.cmeta {padding: 28pt 48pt; color: #4b5563; font-size: 10pt;}
"""

SLIDE_CSS = """
.slide {padding: 34pt 44pt;}
.stitle {color: #1e1b4b; font-size: 26pt; font-weight: bold; margin-bottom: 16pt;}
.bar {background-color: #8135cd; padding: 3pt; margin-bottom: 18pt;}
table.kpi td {border: none; background-color: #f5f3ff; padding: 14pt 10pt; text-align: center;}
.kv1 {font-size: 22pt; font-weight: bold; color: #8135cd;}
.kl {font-size: 10pt; color: #6b7280;}
.kd {font-size: 10pt; color: #059669;}
li {font-size: 15pt; margin-bottom: 8pt;}
td, th {font-size: 12pt;}
"""

ARCH = fitz.Archive(str(FONTS))


def esc(s):
    return html.escape(str(s))


def table_html(rows, cls="", bold_last=False):
    head = "".join(f"<th>{esc(c)}</th>" for c in rows[0])
    body = []
    for i, r in enumerate(rows[1:]):
        last = bold_last and i >= len(rows) - 4 and not r[0]
        cls_attr = ' class="last"' if last else ""
        body.append(f"<tr{cls_attr}>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>")
    return f'<table class="{cls}"><tr>{head}</tr>{"".join(body)}</table>'


def blocks_html(blocks):
    out = []
    for b in blocks:
        k = b[0]
        if k == "h1":
            out.append(f"<h1>{esc(b[1])}</h1>")
        elif k == "h2":
            out.append(f"<h2>{esc(b[1])}</h2>")
        elif k == "p":
            out.append(f"<p>{b[1]}</p>")
        elif k == "note":
            out.append(f'<p class="note">{esc(b[1])}</p>')
        elif k == "draft":
            out.append(f'<p class="draft">{esc(b[1])}</p>')
        elif k == "bullets":
            out.append("<ul>" + "".join(f"<li>{esc(x)}</li>" for x in b[1]) + "</ul>")
        elif k == "table":
            out.append(table_html(b[1], bold_last=True))
        elif k == "kv":
            out.append('<table class="kv">' + "".join(f'<tr><td class="k">{esc(a)}</td><td>{esc(v)}</td></tr>' for a, v in b[1]) + "</table>")
        elif k == "rtl":
            out.append(f'<div dir="rtl"><p class="ar">{esc(b[1])}</p></div>')
        elif k == "rtl_h":
            out.append(f'<div dir="rtl"><h1 class="ar">{esc(b[1])}</h1></div>')
        elif k == "rtl_h2":
            out.append(f'<div dir="rtl"><h2 class="ar">{esc(b[1])}</h2></div>')
        elif k == "rtl_bullets":
            out.append('<div dir="rtl">' + "".join(f'<p class="ar">&#8226; {esc(x)}</p>' for x in b[1]) + "</div>")
        elif k == "sign":
            cells = "".join(f'<td><p class="line"><b>{esc(n)}</b><br/>{esc(r)}</p></td>' for n, r in b[1])
            out.append(f'<table class="sign"><tr>{cells}</tr></table>')
        elif k == "swatches":
            cells = "".join(f'<td class="swatch" style="background-color:{c}; color:{"#1e1b4b" if c in ("#F5F3FF",) else "#ffffff"}">{esc(n)}<br/>{c}<br/>{esc(u)}</td>' for c, n, u in b[1])
            out.append(f"<table><tr>{cells}</tr></table>")
    return "\n".join(out)


def story_pages(writer, html_text, css, rect, where):
    story = fitz.Story(html=html_text, user_css=css, archive=ARCH)
    more = 1
    while more:
        dev = writer.begin_page(rect)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()


def has_arabic(text):
    return any("\u0600" <= ch <= "\u06ff" for ch in text)


def build_pdf(key, spec, out_dir):
    path = out_dir / f"{key}.pdf"
    tmp = io.BytesIO()
    arabic = has_arabic(json.dumps(spec, ensure_ascii=False))
    css = CSS % {"font": "dubai" if arabic else "segoe"}
    writer = fitz.DocumentWriter(tmp)
    if spec.get("slides"):
        rect = fitz.Rect(0, 0, 960, 540)
        cover = f'<div class="cover" style="padding:170pt 60pt 190pt 60pt"><p class="kicker">ORGTIK</p><p class="ctitle">{esc(spec["title"])}</p><p class="csub">{esc(spec["subtitle"])}</p></div>'
        story_pages(writer, cover, css + COVER_CSS, rect, rect)
        for b in spec["blocks"]:
            _, title, parts = b
            inner = [f'<div class="bar"></div><p class="stitle">{esc(title)}</p>']
            for kind, data in parts:
                if kind == "kpi":
                    cells = "".join(f'<td><p class="kv1">{esc(v)}</p><p class="kl">{esc(l)}</p><p class="kd">{esc(d)}</p></td>' for l, v, d in data)
                    inner.append(f'<table class="kpi"><tr>{cells}</tr></table>')
                elif kind == "table":
                    inner.append(table_html(data, bold_last=False))
                elif kind == "bullets":
                    inner.append("<ul>" + "".join(f"<li>{esc(x)}</li>" for x in data) + "</ul>")
            story_pages(writer, f'<div class="slide">{"".join(inner)}</div>', css + SLIDE_CSS, rect, rect)
    else:
        rect = fitz.paper_rect("a4")
        if spec.get("cover"):
            cover = (f'<div class="cover"><p class="kicker">ORGTIK</p><p class="ctitle">{esc(spec["title"])}</p><p class="csub">{esc(spec["subtitle"])}</p></div>'
                     f'<div class="cmeta">Orgtik LLC &#183; {esc(C.ORG_ADDR)}<br/>Confidential - for internal and client use only</div>')
            story_pages(writer, cover, css + COVER_CSS, rect, rect)
        band = f'<div class="band"><b>ORGTIK</b> &#160; {esc(spec["title"])}</div>'
        story_pages(writer, band + blocks_html(spec["blocks"]), css, rect, rect + (48, 44, -48, -54))
    writer.close()

    doc = fitz.open(stream=tmp.getvalue(), filetype="pdf")
    n = doc.page_count
    footer_title = "".join(ch for ch in spec["title"] if ord(ch) < 128).strip(" -") or "Document"
    for i, page in enumerate(doc):
        if i == 0 and (spec.get("cover") or spec.get("slides")):
            continue
        r = page.rect
        page.draw_line((48, r.height - 36), (r.width - 48, r.height - 36), color=(0.85, 0.85, 0.88), width=0.6)
        page.insert_text((48, r.height - 24), f"Orgtik LLC  |  {footer_title}", fontsize=7.5, fontname="helv", color=(0.45, 0.45, 0.5))
        label = f"Page {i + 1} of {n}"
        page.insert_text((r.width - 48 - fitz.get_text_length(label, "helv", 7.5), r.height - 24), label, fontsize=7.5, fontname="helv", color=(0.45, 0.45, 0.5))
    doc.set_metadata({"title": spec["title"], "author": "Orgtik LLC", "creator": "Orgtik", "producer": "Orgtik local-dev"})
    doc.subset_fonts()
    doc.save(str(path), garbage=4, deflate=True)
    doc.close()
    return path


def render_pages(key, pdf_path, max_pages=6, width=900):
    out = ASSETS / "previews" / key
    out.mkdir(parents=True, exist_ok=True)
    doc = pdf_path if isinstance(pdf_path, fitz.Document) else fitz.open(str(pdf_path))
    assets = []
    for i, page in enumerate(doc):
        if i >= max_pages:
            break
        zoom = width / page.rect.width
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        img = img.quantize(colors=128, method=Image.Quantize.MEDIANCUT)
        p = out / f"page-{i + 1}.png"
        img.save(p, optimize=True)
        assets.append({"key": p.name, "file": rel(p), "mimeType": "image/png", "size": p.stat().st_size})
    doc.close()
    return {"kind": "pages", "pageCount": len(assets), "assets": assets}


def preview_from_html(key, html_text, landscape=False, font_family="segoe"):
    tmp = io.BytesIO()
    css = CSS % {"font": font_family}
    writer = fitz.DocumentWriter(tmp)
    rect = fitz.paper_rect("a4-l" if landscape else "a4")
    story_pages(writer, html_text, css, rect, rect + (40, 40, -40, -40))
    writer.close()
    return render_pages(key, fitz.open(stream=tmp.getvalue(), filetype="pdf"), width=1000 if landscape else 900)


# ============================================================== DOCX
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def x(s):
    return html.escape(str(s), quote=False)


def w_run(text, bold=False, rtl=False, size=None, color=None):
    rpr = []
    if bold:
        rpr.append("<w:b/><w:bCs/>")
    if color:
        rpr.append(f'<w:color w:val="{color}"/>')
    if size:
        rpr.append(f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>')
    if rtl:
        rpr.append("<w:rtl/>")
    return f'<w:r><w:rPr>{"".join(rpr)}</w:rPr><w:t xml:space="preserve">{x(text)}</w:t></w:r>'


def w_par(text, style=None, rtl=False, bullet=False, bold=False):
    ppr = []
    if style:
        ppr.append(f'<w:pStyle w:val="{style}"/>')
    if bullet:
        ppr.append('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
    if rtl:
        ppr.append("<w:bidi/>")
    return f'<w:p><w:pPr>{"".join(ppr)}</w:pPr>{w_run(text, bold=bold, rtl=rtl)}</w:p>'


def w_table(rows, rtl=False):
    ncol = len(rows[0])
    colw = 9000 // ncol
    borders = "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="D9D9E3"/>' for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
    out = [f'<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/>{"<w:bidiVisual/>" if rtl else ""}<w:tblBorders>{borders}</w:tblBorders>'
           '<w:tblCellMar><w:top w:w="60" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:bottom w:w="60" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>'
           '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr>']
    out.append("<w:tblGrid>" + "".join(f'<w:gridCol w:w="{colw}"/>' for _ in range(ncol)) + "</w:tblGrid>")
    for i, row in enumerate(rows):
        cells = []
        for c in row:
            shade = '<w:shd w:val="clear" w:color="auto" w:fill="F3E8FF"/>' if i == 0 else ""
            ppr = "<w:bidi/>" if rtl else ""
            cells.append(f'<w:tc><w:tcPr><w:tcW w:w="{colw}" w:type="dxa"/>{shade}</w:tcPr><w:p><w:pPr><w:spacing w:after="0"/>{ppr}</w:pPr>{w_run(c, bold=(i == 0), rtl=rtl)}</w:p></w:tc>')
        out.append("<w:tr>" + "".join(cells) + "</w:tr>")
    out.append("</w:tbl><w:p/>")
    return "".join(out)


DOCX_STYLES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles {W_NS}>
<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Calibri" w:cs="Arial"/><w:sz w:val="22"/><w:szCs w:val="24"/><w:lang w:val="en-US" w:bidi="ar-EG"/></w:rPr></w:rPrDefault>
<w:pPrDefault><w:pPr><w:spacing w:after="140" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>
<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="240"/></w:pPr><w:rPr><w:b/><w:bCs/><w:color w:val="8135CD"/><w:sz w:val="44"/><w:szCs w:val="44"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="360" w:after="120"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:bCs/><w:color w:val="8135CD"/><w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="80"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:b/><w:bCs/><w:color w:val="1E1B4B"/><w:sz w:val="26"/><w:szCs w:val="26"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="60"/><w:ind w:left="720"/></w:pPr></w:style>
<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>
</w:styles>"""

DOCX_NUMBERING = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering {W_NS}>
<w:abstractNum w:abstractNumId="0"><w:multiLevelType w:val="hybridMultilevel"/>
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="&#8226;"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>
</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>"""

DOCX_SETTINGS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:settings {W_NS}><w:defaultTabStop w:val="720"/><w:characterSpacingControl w:val="doNotCompress"/>
<w:compat><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat></w:settings>"""


def core_props(title, creator="Orgtik LLC"):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f"<dc:title>{x(title)}</dc:title><dc:creator>{x(creator)}</dc:creator><cp:lastModifiedBy>{x(creator)}</cp:lastModifiedBy>"
            '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-01T09:00:00Z</dcterms:created><dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-20T09:00:00Z</dcterms:modified>'
            "</cp:coreProperties>")


def app_props(app):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            f"<Application>{app}</Application><Company>Orgtik LLC</Company></Properties>")


def write_zip(path, entries):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 20, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)


def build_docx(key, spec, out_dir):
    rtl_doc = bool(spec.get("rtl"))
    body = []
    for b in spec["blocks"]:
        k = b[0]
        if k == "h1":
            body.append(w_par(b[1], "Heading1"))
        elif k == "h2":
            body.append(w_par(b[1], "Heading2"))
        elif k == "p":
            body.append(w_par(b[1]))
        elif k == "bullets":
            body += [w_par(t, "ListParagraph", bullet=True) for t in b[1]]
        elif k == "table":
            body.append(w_table(b[1], rtl=rtl_doc))
        elif k == "rtl_h":
            body.append(w_par(b[1], "Heading1", rtl=True))
        elif k == "rtl_h2":
            body.append(w_par(b[1], "Heading2", rtl=True))
        elif k == "rtl":
            body.append(w_par(b[1], rtl=True))
        elif k == "rtl_bullets":
            body += [w_par(t, "ListParagraph", rtl=True, bullet=True) for t in b[1]]
    sect = '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1300" w:right="1200" w:bottom="1300" w:left="1200" w:header="708" w:footer="708" w:gutter="0"/>' + ("<w:bidi/>" if rtl_doc else "") + "</w:sectPr>"
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:document {W_NS}><w:body>{"".join(body)}{sect}</w:body></w:document>'
    rels_root = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                 '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
                 '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/></Relationships>')
    rels_doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>'
                '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" Target="settings.xml"/></Relationships>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
              '<Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>'
              '<Override PartName="/word/settings.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
              '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
              '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/></Types>')
    path = out_dir / f"{key}.docx"
    write_zip(path, [("[Content_Types].xml", ctypes), ("_rels/.rels", rels_root), ("word/document.xml", document), ("word/_rels/document.xml.rels", rels_doc),
                     ("word/styles.xml", DOCX_STYLES), ("word/numbering.xml", DOCX_NUMBERING), ("word/settings.xml", DOCX_SETTINGS),
                     ("docProps/core.xml", core_props(spec["title"])), ("docProps/app.xml", app_props("Microsoft Office Word"))])
    blocks = [("h1", spec["title"])] if not rtl_doc and spec["blocks"][0][0] != "h1" else []
    prev = preview_from_html(key, blocks_html(blocks + spec["blocks"]), font_family="dubai" if rtl_doc else "segoe")
    return path, prev


# ============================================================== XLSX
def build_xlsx(key, sheets, out_dir):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="8135CD")
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value)) for c in col if c.value is not None)
            ws.column_dimensions[col[0].column_letter].width = min(42, max(10, width + 2))
            for c in col[1:]:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0.00" if isinstance(c.value, float) else "#,##0"
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22
    wb.properties.creator = "Orgtik LLC"
    wb.properties.title = key
    path = out_dir / f"{key}.xlsx"
    wb.save(path)
    name, header, rows = sheets[0]
    fmt = lambda v: f"{v:,.2f}" if isinstance(v, float) else (f"{v:,}" if isinstance(v, int) else v)  # noqa: E731
    tbl = table_html([header] + [[fmt(v) for v in r] for r in rows])
    prev = preview_from_html(key, f'<div class="band"><b>ORGTIK</b> &#160; Sheet: {esc(name)}</div>{tbl}', landscape=True)
    return path, prev


# ============================================================== PPTX
P_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
GRP = ('<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
       '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>')
THEME = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Orgtik"><a:themeElements>
<a:clrScheme name="Orgtik"><a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1><a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
<a:dk2><a:srgbClr val="1E1B4B"/></a:dk2><a:lt2><a:srgbClr val="F5F3FF"/></a:lt2><a:accent1><a:srgbClr val="8135CD"/></a:accent1><a:accent2><a:srgbClr val="FF692E"/></a:accent2>
<a:accent3><a:srgbClr val="32ABB9"/></a:accent3><a:accent4><a:srgbClr val="0061FF"/></a:accent4><a:accent5><a:srgbClr val="F59E0B"/></a:accent5><a:accent6><a:srgbClr val="10B981"/></a:accent6>
<a:hlink><a:srgbClr val="0061FF"/></a:hlink><a:folHlink><a:srgbClr val="7C3AED"/></a:folHlink></a:clrScheme>
<a:fontScheme name="Orgtik"><a:majorFont><a:latin typeface="Segoe UI"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont><a:minorFont><a:latin typeface="Segoe UI"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont></a:fontScheme>
<a:fmtScheme name="Orgtik"><a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>
<a:lnStyleLst><a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="19050"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>
<a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>
<a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst></a:fmtScheme>
</a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>"""


def p_shape(sid, name, x0, y0, cx, cy, fill=None, paras="", anchor="t"):
    fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:noFill/></a:ln>' if fill else "<a:noFill/>"
    empty_para = '<a:p><a:endParaRPr lang="en-US"/></a:p>'
    txbox = ' txBox="1"' if not fill else ""
    body = (f'<p:txBody><a:bodyPr wrap="square" rtlCol="0" anchor="{anchor}"><a:normAutofit/></a:bodyPr><a:lstStyle/>{paras or empty_para}</p:txBody>')
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr{txbox}/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{x0}" y="{y0}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill_xml}</p:spPr>{body}</p:sp>')


def p_para(text, size, color, bold=False, bullet=False):
    ppr = '<a:pPr marL="457200" indent="-342900"><a:spcBef><a:spcPts val="900"/></a:spcBef><a:buFont typeface="Arial"/><a:buChar char="&#8226;"/></a:pPr>' if bullet else "<a:pPr><a:buNone/></a:pPr>"
    battr = ' b="1"' if bold else ""
    return (f'<a:p>{ppr}<a:r><a:rPr lang="en-US" sz="{size}"{battr} dirty="0"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="Segoe UI"/></a:rPr><a:t>{x(text)}</a:t></a:r></a:p>')


def build_pptx(key, spec, out_dir):
    slides_xml = []
    for i, (title, lines) in enumerate(spec["slides"]):
        if i == 0:
            shapes = [p_shape(2, "Background", 0, 0, 12192000, 6858000, fill="8135CD"),
                      p_shape(3, "Title", 914400, 2286000, 10363200, 1371600, paras=p_para(title, 5400, "FFFFFF", bold=True), anchor="b"),
                      p_shape(4, "Subtitle", 914400, 3810000, 10363200, 1371600, paras="".join(p_para(t, 2400, "EDE9FE") for t in lines))]
        else:
            shapes = [p_shape(2, "Accent", 0, 0, 12192000, 182880, fill="8135CD"),
                      p_shape(3, "Title", 685800, 457200, 10820400, 1143000, paras=p_para(title, 4000, "1E1B4B", bold=True), anchor="b"),
                      p_shape(4, "Body", 685800, 1828800, 10820400, 4389120, paras="".join(p_para(t, 2400, "374151", bullet=True) for t in lines)),
                      p_shape(5, "Footer", 685800, 6309360, 10820400, 365760, paras=p_para(f"Orgtik  |  {spec['title']}  |  {i + 1}", 1100, "9CA3AF"))]
        bg = '<p:bg><p:bgPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'
        slides_xml.append(f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sld {P_NS}><p:cSld>{bg}<p:spTree>{GRP}{"".join(shapes)}</p:spTree></p:cSld>'
                          '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>')
    n = len(slides_xml)
    rel_head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    pres = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:presentation {P_NS} saveSubsetFonts="1"><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
            "<p:sldIdLst>" + "".join(f'<p:sldId id="{256 + i}" r:id="rId{2 + i}"/>' for i in range(n)) + "</p:sldIdLst>"
            '<p:sldSz cx="12192000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>')
    pres_rels = (rel_head + f'<Relationship Id="rId1" Type="{R}slideMaster" Target="slideMasters/slideMaster1.xml"/>'
                 + "".join(f'<Relationship Id="rId{2 + i}" Type="{R}slide" Target="slides/slide{i + 1}.xml"/>' for i in range(n))
                 + f'<Relationship Id="rId{n + 2}" Type="{R}presProps" Target="presProps.xml"/><Relationship Id="rId{n + 3}" Type="{R}viewProps" Target="viewProps.xml"/>'
                 + f'<Relationship Id="rId{n + 4}" Type="{R}theme" Target="theme/theme1.xml"/><Relationship Id="rId{n + 5}" Type="{R}tableStyles" Target="tableStyles.xml"/></Relationships>')
    master = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sldMaster {P_NS}><p:cSld><p:bg><p:bgRef idx="1001"><a:schemeClr val="bg1"/></p:bgRef></p:bg><p:spTree>{GRP}</p:spTree></p:cSld>'
              '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
              '<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst>'
              '<p:txStyles><p:titleStyle><a:lvl1pPr><a:defRPr sz="4400"/></a:lvl1pPr></p:titleStyle><p:bodyStyle><a:lvl1pPr><a:defRPr sz="2400"/></a:lvl1pPr></p:bodyStyle>'
              '<p:otherStyle><a:lvl1pPr><a:defRPr sz="1800"/></a:lvl1pPr></p:otherStyle></p:txStyles></p:sldMaster>')
    master_rels = rel_head + f'<Relationship Id="rId1" Type="{R}slideLayout" Target="../slideLayouts/slideLayout1.xml"/><Relationship Id="rId2" Type="{R}theme" Target="../theme/theme1.xml"/></Relationships>'
    layout = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sldLayout {P_NS} type="blank" preserve="1"><p:cSld name="Blank"><p:spTree>{GRP}</p:spTree></p:cSld>'
              '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>')
    layout_rels = rel_head + f'<Relationship Id="rId1" Type="{R}slideMaster" Target="../slideMasters/slideMaster1.xml"/></Relationships>'
    slide_rels = rel_head + f'<Relationship Id="rId1" Type="{R}slideLayout" Target="../slideLayouts/slideLayout1.xml"/></Relationships>'
    root_rels = (rel_head + f'<Relationship Id="rId1" Type="{R}officeDocument" Target="ppt/presentation.xml"/>'
                 '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
                 f'<Relationship Id="rId3" Type="{R}extended-properties" Target="docProps/app.xml"/></Relationships>')
    PT = "application/vnd.openxmlformats-officedocument.presentationml."
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
              f'<Override PartName="/ppt/presentation.xml" ContentType="{PT}presentation.main+xml"/>'
              f'<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="{PT}slideMaster+xml"/>'
              f'<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="{PT}slideLayout+xml"/>'
              + "".join(f'<Override PartName="/ppt/slides/slide{i + 1}.xml" ContentType="{PT}slide+xml"/>' for i in range(n))
              + '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
              f'<Override PartName="/ppt/presProps.xml" ContentType="{PT}presProps+xml"/><Override PartName="/ppt/viewProps.xml" ContentType="{PT}viewProps+xml"/>'
              f'<Override PartName="/ppt/tableStyles.xml" ContentType="{PT}tableStyles+xml"/>'
              '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
              '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/></Types>')
    entries = [("[Content_Types].xml", ctypes), ("_rels/.rels", root_rels), ("ppt/presentation.xml", pres), ("ppt/_rels/presentation.xml.rels", pres_rels),
               ("ppt/slideMasters/slideMaster1.xml", master), ("ppt/slideMasters/_rels/slideMaster1.xml.rels", master_rels),
               ("ppt/slideLayouts/slideLayout1.xml", layout), ("ppt/slideLayouts/_rels/slideLayout1.xml.rels", layout_rels),
               ("ppt/theme/theme1.xml", THEME),
               ("ppt/presProps.xml", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:presentationPr {P_NS}/>'),
               ("ppt/viewProps.xml", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:viewPr {P_NS}><p:normalViewPr><p:restoredLeft sz="15620"/><p:restoredTop sz="94660"/></p:normalViewPr><p:gridSpacing cx="76200" cy="76200"/></p:viewPr>'),
               ("ppt/tableStyles.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" def="{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"/>'),
               ("docProps/core.xml", core_props(spec["title"])), ("docProps/app.xml", app_props("Microsoft Office PowerPoint"))]
    for i, s in enumerate(slides_xml):
        entries.append((f"ppt/slides/slide{i + 1}.xml", s))
        entries.append((f"ppt/slides/_rels/slide{i + 1}.xml.rels", slide_rels))
    path = out_dir / f"{key}.pptx"
    write_zip(path, entries)
    # preview: same slides rendered as landscape pages
    tmp = io.BytesIO()
    writer = fitz.DocumentWriter(tmp)
    rect = fitz.Rect(0, 0, 960, 540)
    css = CSS % {"font": "segoe"} + COVER_CSS + SLIDE_CSS
    for i, (title, lines) in enumerate(spec["slides"]):
        if i == 0:
            h = f'<div class="cover" style="padding:170pt 60pt 190pt 60pt"><p class="ctitle">{esc(title)}</p>' + "".join(f'<p class="csub">{esc(t)}</p>' for t in lines) + "</div>"
        else:
            h = f'<div class="slide"><div class="bar"></div><p class="stitle">{esc(title)}</p><ul>' + "".join(f"<li>{esc(t)}</li>" for t in lines) + "</ul></div>"
        story_pages(writer, h, css, rect, rect)
    writer.close()
    prev = render_pages(key, fitz.open(stream=tmp.getvalue(), filetype="pdf"), width=960)
    return path, prev


# ============================================================== images (Pillow)
def hexc(h, a=255):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4)) + (a,)


def gradient(size, c1, c2, vertical=True, diagonal=False):
    w, h = size
    base = Image.new("RGB", (w, h))
    a, b = hexc(c1)[:3], hexc(c2)[:3]
    px = base.load()
    for yy in range(h):
        for xx in range(w):
            t = ((xx / w + yy / h) / 2) if diagonal else (yy / h if vertical else xx / w)
            px[xx, yy] = tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3))
    return base


def fast_gradient(size, c1, c2, vertical=True):
    w, h = size
    strip = Image.new("RGB", (1, 256) if vertical else (256, 1))
    a, b = hexc(c1)[:3], hexc(c2)[:3]
    for i in range(256):
        t = i / 255
        strip.putpixel((0, i) if vertical else (i, 0), tuple(int(a[k] + (b[k] - a[k]) * t) for k in range(3)))
    return strip.resize((w, h), Image.BILINEAR)


class Canvas:
    """2x supersampled drawing surface -> downscaled on save for smooth edges."""

    def __init__(self, w, h, bg="#ffffff", grad=None, ss=2):
        self.w, self.h, self.ss = w, h, ss
        if grad:
            self.img = fast_gradient((w * ss, h * ss), grad[0], grad[1], vertical=grad[2] if len(grad) > 2 else True).convert("RGB")
        else:
            self.img = Image.new("RGB", (w * ss, h * ss), hexc(bg)[:3])
        self.d = ImageDraw.Draw(self.img, "RGBA")

    def s(self, *v):
        return [int(q * self.ss) for q in v]

    def rect(self, x0, y0, x1, y1, fill, r=0, outline=None, width=1):
        self.d.rounded_rectangle(self.s(x0, y0, x1, y1), radius=int(r * self.ss), fill=hexc(fill) if isinstance(fill, str) else fill,
                                 outline=hexc(outline) if outline else None, width=int(width * self.ss))

    def ellipse(self, x0, y0, x1, y1, fill, outline=None, width=1):
        self.d.ellipse(self.s(x0, y0, x1, y1), fill=hexc(fill) if isinstance(fill, str) else fill, outline=hexc(outline) if outline else None, width=int(width * self.ss))

    def poly(self, pts, fill):
        self.d.polygon([(int(px * self.ss), int(py * self.ss)) for px, py in pts], fill=hexc(fill) if isinstance(fill, str) else fill)

    def line(self, pts, fill, width=2):
        self.d.line([(int(px * self.ss), int(py * self.ss)) for px, py in pts], fill=hexc(fill), width=int(width * self.ss), joint="curve")

    def text(self, xy, txt, size, fill="#111827", bold=False, anchor="la", face=None):
        f = font(face or ("segoeuib.ttf" if bold else "segoeui.ttf"), int(size * self.ss))
        self.d.text((xy[0] * self.ss, xy[1] * self.ss), txt, font=f, fill=hexc(fill) if isinstance(fill, str) else fill, anchor=anchor)

    def shadow(self, x0, y0, x1, y1, r=16, blur=18, alpha=70):
        layer = Image.new("RGBA", self.img.size, (0, 0, 0, 0))
        ImageDraw.Draw(layer).rounded_rectangle(self.s(x0, y0, x1, y1), radius=int(r * self.ss), fill=(20, 10, 40, alpha))
        layer = layer.filter(ImageFilter.GaussianBlur(blur * self.ss / 2))
        self.img.paste(layer, (0, 0), layer)
        self.d = ImageDraw.Draw(self.img, "RGBA")

    def final(self):
        return self.img.resize((self.w, self.h), Image.LANCZOS).convert("RGB")


def save_img(key, img, fmt, out_dir):
    ext = "jpg" if fmt == "JPEG" else "png"
    path = out_dir / f"{key}.{ext}"
    if fmt == "JPEG":
        img.save(path, "JPEG", quality=86, optimize=True, progressive=True)
    else:
        img.save(path, "PNG", optimize=True)
    register(key, path, "image/jpeg" if fmt == "JPEG" else "image/png", {"kind": "image"})


def img_insta_post():
    c = Canvas(800, 800, grad=("#9B4DE0", "#4C1D95"))
    c.ellipse(560, -120, 980, 300, (255, 255, 255, 28))
    c.ellipse(-160, 560, 260, 980, (255, 255, 255, 22))
    c.rect(56, 64, 170, 104, "#FF692E", r=20)
    c.text((113, 84), "NEW", 20, "#ffffff", bold=True, anchor="mm")
    c.text((56, 140), "Orgtik 3.0", 78, "#ffffff", bold=True)
    c.text((58, 240), "Everything your team needs.", 30, "#EDE9FE")
    c.text((58, 280), "Connected.", 30, "#EDE9FE", bold=True)
    c.shadow(430, 330, 700, 820, r=40, blur=30, alpha=110)
    c.rect(430, 330, 700, 840, "#111827", r=40)
    c.rect(444, 346, 686, 840, "#F7F5FC", r=30)
    c.rect(444, 346, 686, 420, "#8135CD", r=30)
    c.rect(444, 390, 686, 420, "#8135CD")
    c.text((462, 372), "Good morning, Salma", 15, "#ffffff", bold=True)
    for i, (label, val, col) in enumerate([("Tasks due", "5", "#FF692E"), ("Shared files", "12", "#32ABB9"), ("Deals", "3", "#0061FF")]):
        y = 440 + i * 86
        c.rect(460, y, 670, y + 72, "#ffffff", r=14)
        c.ellipse(474, y + 18, 510, y + 54, col)
        c.text((524, y + 18), label, 14, "#6B7280")
        c.text((524, y + 36), val, 22, "#111827", bold=True)
    c.text((56, 700), "Tasks  ·  Files  ·  Sales  ·  HR  ·  Marketing", 22, "#DDD6FE")
    c.text((56, 740), "orgtik.test/3.0", 20, "#C4B5FD", bold=True)
    return c.final()


def img_linkedin():
    c = Canvas(1200, 300, grad=("#F5F3FF", "#EDE9FE", False))
    for i, (x0, y0, r, a) in enumerate([(930, -60, 260, 60), (1080, 140, 180, 90), (820, 170, 120, 40)]):
        c.ellipse(x0, y0, x0 + r, y0 + r, (129, 53, 205, a))
    c.text((70, 70), "Run your business on one", 44, "#1E1B4B", bold=True)
    c.text((70, 124), "connected suite.", 44, "#8135CD", bold=True)
    c.text((72, 196), "Tasks · Files · Sales · HR · Marketing  —  Cairo · Riyadh · Dubai", 20, "#4B5563")
    c.rect(900, 110, 1110, 190, "#ffffff", r=18)
    c.text((1005, 150), "orgtik", 40, "#8135CD", bold=True, anchor="mm")
    return c.final()


def img_booth():
    c = Canvas(1200, 800, grad=("#E5E7EB", "#F9FAFB"))
    c.poly([(80, 640), (1120, 640), (1000, 760), (200, 760)], "#D1D5DB")  # floor
    c.poly([(200, 180), (900, 120), (900, 600), (200, 640)], "#8135CD")  # back wall
    c.poly([(900, 120), (1080, 200), (1080, 600), (900, 600)], "#6D28D9")  # side wall
    c.poly([(200, 180), (900, 120), (900, 160), (200, 220)], "#5B21B6")  # header strip
    c.text((550, 250), "ORGTIK", 86, "#ffffff", bold=True, anchor="mm")
    c.text((550, 318), "Your business, connected.", 26, "#EDE9FE", anchor="mm")
    for i in range(2):
        x0 = 280 + i * 330
        c.rect(x0, 380, x0 + 240, 520, "#111827", r=6)
        c.rect(x0 + 10, 390, x0 + 230, 510, "#1E1B4B", r=4)
        for j in range(5):
            c.rect(x0 + 24 + j * 40, 490 - (j + 2) * 12, x0 + 50 + j * 40, 500, "#32ABB9" if j % 2 else "#FF692E")
    c.poly([(420, 560), (720, 560), (760, 700), (380, 700)], "#F9FAFB")  # counter
    c.poly([(380, 700), (760, 700), (760, 716), (380, 716)], "#9CA3AF")
    c.text((570, 632), "B-214", 30, "#8135CD", bold=True, anchor="mm")
    for px in (240, 960):
        c.rect(px - 18, 600, px + 18, 680, "#78350F", r=6)
        for dx, dy, r in ((-40, -70, 60), (10, -90, 60), (-15, -120, 55)):
            c.ellipse(px + dx, 600 + dy, px + dx + r, 600 + dy + r, "#16A34A")
    c.text((40, 40), "Riyadh Expo 2026 · Booth B-214 · Final render v3", 22, "#374151", bold=True)
    return c.final()


def phone_frame(c, title, dark_header=True):
    c.rect(0, 0, 430, 60, "#8135CD" if dark_header else "#ffffff")
    c.text((24, 20), "9:41", 17, "#ffffff" if dark_header else "#111827", bold=True)
    c.text((406, 20), "5G  100%", 15, "#ffffff" if dark_header else "#111827", anchor="ra")
    c.rect(0, 60, 430, 130, "#8135CD")
    c.text((24, 80), title, 26, "#ffffff", bold=True)


def img_app_home():
    c = Canvas(430, 932, bg="#F7F5FC")
    phone_frame(c, "Hi Nour 👋".replace("👋", ""))
    c.rect(20, 150, 410, 196, "#ffffff", r=23)
    c.text((48, 162), "Search juices, dairy, snacks…", 16, "#9CA3AF")
    c.text((20, 220), "Fresh this week", 20, "#111827", bold=True)
    for i, (name, price, col) in enumerate([("Mango 1L", "EGP 42", "#F59E0B"), ("Guava 1L", "EGP 39", "#EC4899"), ("Orange 1L", "EGP 36", "#F97316")]):
        x0 = 20 + i * 136
        c.rect(x0, 256, x0 + 124, 430, "#ffffff", r=18)
        c.rect(x0 + 40, 276, x0 + 84, 380, col, r=10)
        c.rect(x0 + 52, 266, x0 + 72, 280, "#374151", r=4)
        c.text((x0 + 14, 390), name, 14, "#111827", bold=True)
        c.text((x0 + 14, 408), price, 13, "#8135CD", bold=True)
    c.text((20, 460), "Quick re-order", 20, "#111827", bold=True)
    for i, (t, sub) in enumerate([("Family Pack 4 x 1L", "Last ordered 12 Sep"), ("Nile Laban 6 x 200ml", "Last ordered 3 Sep"), ("Breakfast bundle", "Last ordered 28 Aug")]):
        y = 500 + i * 92
        c.rect(20, y, 410, y + 80, "#ffffff", r=16)
        c.rect(36, y + 14, 88, y + 66, "#EDE9FE", r=12)
        c.text((104, y + 18), t, 16, "#111827", bold=True)
        c.text((104, y + 44), sub, 13, "#6B7280")
        c.rect(330, y + 24, 394, y + 56, "#8135CD", r=16)
        c.text((362, y + 40), "Add", 13, "#ffffff", bold=True, anchor="mm")
    c.rect(0, 852, 430, 932, "#ffffff")
    for i, t in enumerate(["Home", "Search", "Orders", "Profile"]):
        c.text((54 + i * 107, 890), t, 13, "#8135CD" if i == 0 else "#9CA3AF", bold=i == 0, anchor="mm")
    return c.final()


def img_app_checkout():
    c = Canvas(430, 932, bg="#F7F5FC")
    phone_frame(c, "Checkout")
    c.rect(20, 150, 410, 330, "#ffffff", r=18)
    c.text((40, 166), "Your order", 17, "#111827", bold=True)
    for i, (t, p) in enumerate([("Mango 1L  x 4", "EGP 168.00"), ("Family Pack 4 x 1L  x 1", "EGP 149.00"), ("Nile Laban 6 x 200ml  x 2", "EGP 132.00")]):
        c.text((40, 204 + i * 36), t, 15, "#374151")
        c.text((390, 204 + i * 36), p, 15, "#111827", bold=True, anchor="ra")
    c.rect(20, 346, 410, 460, "#ffffff", r=18)
    c.text((40, 362), "Deliver to", 17, "#111827", bold=True)
    c.text((40, 394), "Home · 12 El Nasr Rd, Nasr City, Cairo", 14, "#374151")
    c.text((40, 420), "Today 18:00 - 20:00", 14, "#8135CD", bold=True)
    c.rect(20, 476, 410, 600, "#ffffff", r=18)
    c.text((40, 492), "Payment", 17, "#111827", bold=True)
    for i, (t, on) in enumerate([("Card •••• 4821", True), ("Cash on delivery", False)]):
        y = 528 + i * 34
        c.ellipse(40, y, 60, y + 20, "#8135CD" if on else "#ffffff", outline="#8135CD", width=2)
        c.text((74, y), t, 15, "#374151")
    c.rect(20, 616, 410, 760, "#ffffff", r=18)
    for i, (t, p) in enumerate([("Subtotal", "EGP 449.00"), ("Delivery", "EGP 15.00"), ("VAT 14% (incl.)", "EGP 55.14"), ("Total", "EGP 464.00")]):
        c.text((40, 632 + i * 30), t, 15 if i < 3 else 17, "#6B7280" if i < 3 else "#111827", bold=i == 3)
        c.text((390, 632 + i * 30), p, 15 if i < 3 else 17, "#111827", bold=i == 3, anchor="ra")
    c.rect(20, 790, 410, 850, "#8135CD", r=30)
    c.text((215, 820), "Place order · EGP 464.00", 18, "#ffffff", bold=True, anchor="mm")
    return c.final()


def img_dashboard_dark():
    c = Canvas(1280, 800, bg="#0F0B1E")
    c.rect(0, 0, 230, 800, "#171229")
    c.text((30, 34), "orgtik", 30, "#A78BFA", bold=True)
    for i, t in enumerate(["Overview", "Projects", "Files", "Sales", "Reports", "Settings"]):
        y = 110 + i * 48
        if i == 0:
            c.rect(16, y - 10, 214, y + 30, "#2E1F57", r=10)
        c.text((40, y), t, 17, "#ffffff" if i == 0 else "#9CA3AF", bold=i == 0)
    c.text((270, 34), "Overview", 30, "#F9FAFB", bold=True)
    c.text((270, 78), "Q3 2026 · all regions", 16, "#9CA3AF")
    for i, (label, val, delta, col) in enumerate([("Revenue", "EGP 21.4M", "+18%", "#A78BFA"), ("Active users", "18,420", "+6%", "#32ABB9"),
                                                  ("Open deals", "146", "+12", "#FF692E"), ("Churn", "1.1%", "-0.4", "#34D399")]):
        x0 = 270 + i * 245
        c.rect(x0, 120, x0 + 225, 230, "#1B1533", r=16)
        c.text((x0 + 20, 138), label, 15, "#9CA3AF")
        c.text((x0 + 20, 164), val, 30, "#F9FAFB", bold=True)
        c.text((x0 + 20, 204), delta + " vs Q2", 14, col, bold=True)
    c.rect(270, 250, 880, 560, "#1B1533", r=16)
    c.text((292, 268), "Revenue trend (EGP M)", 17, "#E5E7EB", bold=True)
    pts = [(310 + i * 52, 520 - v * 15) for i, v in enumerate([6, 7, 6.5, 8, 9, 8.5, 10, 11, 12, 11.5, 13.5, 15])]
    for gx in range(5):
        c.line([(300, 330 + gx * 45), (860, 330 + gx * 45)], "#2A2346", width=1)
    c.line(pts, "#A78BFA", width=4)
    for p in pts:
        c.ellipse(p[0] - 5, p[1] - 5, p[0] + 5, p[1] + 5, "#A78BFA")
    c.rect(900, 250, 1250, 560, "#1B1533", r=16)
    c.text((922, 268), "Deals by stage", 17, "#E5E7EB", bold=True)
    for i, (t, v, col) in enumerate([("Qualified", 0.9, "#32ABB9"), ("Discovery", 0.7, "#0061FF"), ("Proposal", 0.55, "#A78BFA"), ("Negotiation", 0.35, "#FF692E"), ("Won", 0.25, "#34D399")]):
        y = 318 + i * 46
        c.text((922, y), t, 14, "#9CA3AF")
        c.rect(1030, y + 2, 1030 + int(200 * v), y + 22, col, r=6)
    c.rect(270, 580, 1250, 770, "#1B1533", r=16)
    c.text((292, 598), "Recent activity", 17, "#E5E7EB", bold=True)
    for i, t in enumerate(["Omar Farouk moved Riyadh Logistics to Negotiation", "Laila Mostafa uploaded Dashboard - Dark Mode.png",
                           "Hany Samir marked INV-2026-0142 as paid"]):
        y = 638 + i * 42
        c.ellipse(292, y, 316, y + 24, ["#A78BFA", "#32ABB9", "#34D399"][i])
        c.text((332, y + 1), t, 15, "#D1D5DB")
    return c.final()


def img_screenshot():
    c = Canvas(1280, 720, bg="#E5E7EB")
    c.rect(20, 20, 1260, 700, "#ffffff", r=12)
    c.rect(20, 20, 1260, 60, "#F3F4F6", r=12)
    c.rect(20, 48, 1260, 60, "#F3F4F6")
    for i, col in enumerate(["#EF4444", "#F59E0B", "#22C55E"]):
        c.ellipse(40 + i * 24, 32, 56 + i * 24, 48, col)
    c.rect(400, 30, 880, 52, "#ffffff", r=11)
    c.text((640, 41), "localhost:3001/en/settings", 14, "#6B7280", anchor="mm")
    c.rect(20, 60, 260, 700, "#F7F9FB")
    for i, t in enumerate(["Profile", "Appearance", "Notifications", "Language", "Security"]):
        y = 96 + i * 44
        if i == 1:
            c.rect(36, y - 10, 244, y + 26, "#EDE9FE", r=8)
        c.text((56, y), t, 16, "#8135CD" if i == 1 else "#374151", bold=i == 1)
    c.text((300, 90), "Appearance", 30, "#111827", bold=True)
    c.text((300, 134), "Choose how Orgtik looks on this device.", 16, "#6B7280")
    for i, (t, on) in enumerate([("Dark mode", True), ("Compact sidebar", False), ("Reduce motion", False), ("Right-to-left layout (Arabic)", True)]):
        y = 190 + i * 70
        c.rect(300, y, 1220, y + 56, "#ffffff", r=10, outline="#E5E7EB")
        c.text((324, y + 16), t, 17, "#111827")
        c.rect(1140, y + 14, 1196, y + 42, "#8135CD" if on else "#D1D5DB", r=14)
        cx = 1182 if on else 1154
        c.ellipse(cx - 11, y + 17, cx + 11, y + 39, "#ffffff")
    c.text((300, 490), "Language", 18, "#111827", bold=True)
    c.rect(300, 520, 700, 568, "#ffffff", r=10, outline="#D1D5DB")
    c.text((320, 534), "English (United States)", 16, "#111827")
    c.rect(1040, 620, 1220, 668, "#8135CD", r=10)
    c.text((1130, 644), "Save changes", 16, "#ffffff", bold=True, anchor="mm")
    return c.final()


def img_summer_banner():
    c = Canvas(1200, 628, grad=("#FDBA74", "#F97316"))
    c.ellipse(860, -80, 1260, 320, (255, 237, 150, 200))
    c.ellipse(900, -40, 1220, 280, (255, 244, 180, 255))
    c.text((70, 150), "SUMMER SALE 2025", 64, "#ffffff", bold=True)
    c.text((74, 236), "Up to 30% off annual plans", 40, "#FFF7ED")
    c.rect(74, 320, 380, 392, "#7C2D12", r=36)
    c.text((227, 356), "Offer ends 31 Aug", 24, "#ffffff", bold=True, anchor="mm")
    c.text((74, 540), "orgtik · Tasks · Files · Sales · HR", 24, "#FFEDD5", bold=True)
    return c.final()


def img_old_logo():
    c = Canvas(512, 512, bg="#F0FDFA")
    c.rect(96, 76, 416, 396, "#0F766E", r=80)
    c.ellipse(166, 146, 346, 326, None, outline="#ffffff", width=26)
    c.ellipse(230, 210, 282, 262, "#5EEAD4")
    c.text((256, 440), "orgtik  ·  concept v1 (2024)", 22, "#115E59", bold=True, anchor="mm")
    return c.final()


def bottle(c, cx, top, w, h, liquid, label_text, sub):
    c.rect(cx - w * 0.18, top, cx + w * 0.18, top + h * 0.08, "#1F2937", r=6)
    c.rect(cx - w * 0.28, top + h * 0.07, cx + w * 0.28, top + h * 0.2, (255, 255, 255, 170), r=12)
    c.rect(cx - w / 2, top + h * 0.17, cx + w / 2, top + h, (255, 255, 255, 150), r=int(w * 0.16))
    c.rect(cx - w / 2 + 8, top + h * 0.3, cx + w / 2 - 8, top + h - 8, liquid, r=int(w * 0.14))
    c.rect(cx - w / 2 + 8, top + h * 0.45, cx + w / 2 - 8, top + h * 0.72, "#ffffff", r=10)
    c.text((cx, top + h * 0.53), "NILE JUICE", int(w * 0.13), "#1E3A8A", bold=True, anchor="mm")
    c.text((cx, top + h * 0.63), label_text, int(w * 0.11), liquid if isinstance(liquid, str) else "#92400E", bold=True, anchor="mm")
    c.text((cx, top + h * 0.69), sub, int(w * 0.07), "#6B7280", anchor="mm")


def img_mango():
    c = Canvas(900, 900, grad=("#FFF7ED", "#FED7AA"))
    c.ellipse(260, 770, 640, 830, (120, 53, 15, 60))
    bottle(c, 450, 120, 300, 680, "#F59E0B", "MANGO", "100% natural · 1 L")
    for x0, y0 in ((110, 600), (690, 640)):
        c.ellipse(x0, y0, x0 + 150, y0 + 120, "#FBBF24")
        c.ellipse(x0 + 20, y0 + 10, x0 + 130, y0 + 90, "#FCD34D")
        c.poly([(x0 + 70, y0 + 6), (x0 + 110, y0 - 40), (x0 + 124, y0 - 8)], "#16A34A")
    return c.final()


def img_family_pack():
    c = Canvas(1000, 700, grad=("#F0F9FF", "#DBEAFE"))
    c.rect(90, 150, 910, 600, "#1E3A8A", r=30)
    c.text((500, 110), "NILE JUICE FAMILY PACK · 4 x 1 L", 34, "#1E3A8A", bold=True, anchor="mm")
    for i, (col, name) in enumerate([("#F59E0B", "MANGO"), ("#EC4899", "GUAVA"), ("#F97316", "ORANGE"), ("#84CC16", "APPLE")]):
        cx = 205 + i * 197
        c.rect(cx - 70, 200, cx + 70, 560, "#ffffff", r=18)
        c.rect(cx - 70, 200, cx + 70, 250, col, r=18)
        c.rect(cx - 70, 232, cx + 70, 250, col)
        c.ellipse(cx - 46, 300, cx + 46, 392, col)
        c.text((cx, 430), name, 22, col, bold=True, anchor="mm")
        c.text((cx, 470), "1 L", 20, "#6B7280", anchor="mm")
    c.text((500, 650), "Save 15% vs single bottles", 24, "#1D4ED8", bold=True, anchor="mm")
    return c.final()


def img_iftar():
    c = Canvas(1200, 800, grad=("#0B1026", "#312E81"))
    rnd = random.Random(7)
    for _ in range(70):
        sx, sy, r = rnd.randint(0, 1200), rnd.randint(0, 380), rnd.choice((1, 1, 2, 2, 3))
        c.ellipse(sx, sy, sx + r, sy + r, "#FDE68A")
    c.ellipse(900, 70, 1040, 210, "#FDE68A")
    c.ellipse(935, 50, 1075, 190, "#121536")
    for i, lx in enumerate((150, 330, 870, 1050)):
        ly = 120 + (i % 2) * 50
        c.line([(lx, 0), (lx, ly)], "#FBBF24", width=2)
        c.poly([(lx - 30, ly + 20), (lx + 30, ly + 20), (lx + 22, ly + 110), (lx - 22, ly + 110)], "#F59E0B")
        c.rect(lx - 18, ly + 36, lx + 18, ly + 96, "#FEF3C7", r=6)
        c.poly([(lx - 24, ly + 20), (lx + 24, ly + 20), (lx, ly)], "#B45309")
    c.rect(80, 560, 1120, 640, "#7C2D12", r=16)
    c.rect(60, 540, 1140, 575, "#F5F5F4", r=12)
    for i in range(9):
        px = 130 + i * 118
        c.ellipse(px - 42, 520, px + 42, 560, "#ffffff")
        c.ellipse(px - 26, 527, px + 26, 552, ["#D97706", "#16A34A", "#DC2626"][i % 3])
    c.text((600, 690), "Orgtik Team Iftar 2026 · Cairo", 38, "#FDE68A", bold=True, anchor="mm")
    c.text((600, 740), "140 colleagues · one table", 22, "#C7D2FE", anchor="mm")
    return c.final()


# ============================================================== icons / zips / misc
ICON_PATHS = {
    "home": '<path d="M3 11l9-8 9 8v9a1 1 0 0 1-1 1h-5v-6h-6v6H4a1 1 0 0 1-1-1z"/>',
    "folder": '<path d="M3 6a2 2 0 0 1 2-2h4l2 2h8a2 2 0 0 1 2 2v10a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2z"/>',
    "file": '<path d="M6 2h8l6 6v12a2 2 0 0 1-2 2H6a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2z"/><path d="M14 2v6h6"/>',
    "upload": '<path d="M12 16V4M6 10l6-6 6 6M4 20h16"/>',
    "download": '<path d="M12 4v12M6 10l6 6 6-6M4 20h16"/>',
    "share": '<circle cx="18" cy="5" r="3"/><circle cx="6" cy="12" r="3"/><circle cx="18" cy="19" r="3"/><path d="M8.6 13.5l6.8 4M15.4 6.5l-6.8 4"/>',
    "search": '<circle cx="11" cy="11" r="7"/><path d="M20 20l-4-4"/>',
    "user": '<circle cx="12" cy="8" r="4"/><path d="M4 21a8 8 0 0 1 16 0"/>',
    "settings": '<circle cx="12" cy="12" r="3"/><path d="M12 2v3M12 19v3M2 12h3M19 12h3M4.9 4.9l2.1 2.1M17 17l2.1 2.1M4.9 19.1L7 17M17 7l2.1-2.1"/>',
    "bell": '<path d="M6 16V11a6 6 0 0 1 12 0v5l2 2H4z"/><path d="M10 20a2 2 0 0 0 4 0"/>',
    "calendar": '<rect x="3" y="5" width="18" height="16" rx="2"/><path d="M3 10h18M8 3v4M16 3v4"/>',
    "chart": '<path d="M4 20V10M10 20V4M16 20v-8M22 20H2"/>',
}


def icon_svg(body, color="#8135CD"):
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="{color}" '
            f'stroke-width="2" stroke-linecap="round" stroke-linejoin="round">{body}</svg>\n')


NILE_LOGO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" width="480" height="160" viewBox="0 0 480 160">
<rect width="480" height="160" rx="24" fill="%(bg)s"/>
<path d="M40 104c30-26 60-26 90 0s60 26 90 0" fill="none" stroke="#38BDF8" stroke-width="10" stroke-linecap="round"/>
<path d="M40 72c30-26 60-26 90 0s60 26 90 0" fill="none" stroke="%(fg)s" stroke-width="10" stroke-linecap="round"/>
<text x="250" y="96" font-family="Segoe UI, Arial, sans-serif" font-size="44" font-weight="700" fill="%(fg)s">Nile Foods</text>
</svg>
"""


def build_misc(out_dir):
    # Icon set zip
    entries = [(f"orgtik-icons-v2/{n}.svg", icon_svg(b)) for n, b in ICON_PATHS.items()]
    entries.append(("orgtik-icons-v2/README.txt", "Orgtik icon set v2\n24 x 24 grid, 2 px stroke, round caps.\nColour: #8135CD (use currentColor in code).\nDesigner: Laila Mostafa\n"))
    p = out_dir / "zip_icons.zip"
    write_zip(p, entries)
    register("zip_icons", p, "application/zip")
    # Nile Foods logo pack (uploaded by the client)
    buf = io.BytesIO()
    c = Canvas(480, 160, bg="#ffffff")
    c.line([(40, 104), (85, 80), (130, 104), (175, 128), (220, 104)], "#38BDF8", width=10)
    c.line([(40, 72), (85, 48), (130, 72), (175, 96), (220, 72)], "#1E3A8A", width=10)
    c.text((250, 58), "Nile Foods", 44, "#1E3A8A", bold=True)
    c.final().save(buf, "PNG", optimize=True)
    p = out_dir / "zip_nile_logo.zip"
    write_zip(p, [("Nile Foods Logo Pack/nile-foods-logo.svg", NILE_LOGO_SVG % {"bg": "#FFFFFF", "fg": "#1E3A8A"}),
                  ("Nile Foods Logo Pack/nile-foods-logo-white.svg", NILE_LOGO_SVG % {"bg": "#1E3A8A", "fg": "#FFFFFF"}),
                  ("Nile Foods Logo Pack/nile-foods-logo.png", buf.getvalue()),
                  ("Nile Foods Logo Pack/brand-colours.txt", "Nile Blue  #1E3A8A\nRiver Sky  #38BDF8\nMango      #F59E0B\nGuava      #EC4899\n")])
    register("zip_nile_logo", p, "application/zip")
    # Figma file (binary blob)
    rnd = random.Random(42)
    p = out_dir / "fig_components.fig"
    p.write_bytes(b"fig-kiwi\x10\x00\x00\x00" + bytes(rnd.getrandbits(8) for _ in range(24000)))
    register("fig_components", p, "application/octet-stream")
    # Copies of the app's own brand assets
    for key, src, mime in (("svg_logo_full", PUBLIC / "images" / "logo" / "full-logo.svg", "image/svg+xml"),
                           ("png_logo_icon", PUBLIC / "images" / "logo" / "icon-logo.png", "image/png")):
        dst = out_dir / (key + src.suffix)
        shutil.copyfile(src, dst)
        register(key, dst, mime, {"kind": "image"})


# ============================================================== videos
VW, VH, FPS, GOP = 256, 144, 10, 30


def vid_ramadan(i):
    c = Canvas(VW, VH, grad=("#0B1026", "#3B1D6E"))
    rnd = random.Random(3)
    for k in range(18):
        sx, sy = rnd.randint(4, 250), rnd.randint(4, 70)
        on = (i // 3 + k) % 4 != 0
        r = 1.6 if on else 0.8
        c.ellipse(sx - r, sy - r, sx + r, sy + r, "#FDE68A" if on else "#8B83B8")
    c.ellipse(196, 14, 236, 54, "#FDE68A")
    c.ellipse(206, 8, 246, 48, "#141234")
    t = min(1.0, max(0.0, (i - 6) / 12))
    if t > 0:
        y = 70 + (1 - t) * 14
        col = tuple(int(q * t + 30 * (1 - t)) for q in (253, 230, 138)) + (255,)
        c.text((128, y), "Ramadan Kareem", 22, col, bold=True, anchor="mm")
    if i >= 22:
        c.text((128, 100), "Orgtik  ·  2027", 12, "#C4B5FD", anchor="mm")
    w = min(1.0, max(0.0, (i - 24) / 20))
    if w > 0:
        c.rect(64, 116, 64 + 128 * w, 119, "#F59E0B", r=1)
    return c.final()


def vid_walkthrough(i):
    c = Canvas(VW, VH, bg="#F4F1FB")
    c.text((12, 12), "Mobile App v3", 11, "#1E1B4B", bold=True)
    c.text((12, 28), ["Browse products", "View details", "Checkout"][min(2, i // 20)], 9, "#6B7280")
    c.rect(96, 6, 160, 140, "#111827", r=10)
    c.rect(100, 10, 156, 136, "#FFFFFF", r=8)
    c.rect(100, 10, 156, 28, "#8135CD", r=8)
    c.rect(100, 22, 156, 28, "#8135CD")
    screen = min(2, i // 20)
    if screen == 0:
        for k in range(4):
            y = 34 + k * 25
            c.rect(104, y, 152, y + 21, "#F3E8FF", r=4)
            c.rect(107, y + 3, 121, y + 18, ["#F59E0B", "#EC4899", "#F97316", "#84CC16"][k], r=3)
            c.rect(125, y + 6, 148, y + 9, "#9CA3AF", r=1)
    elif screen == 1:
        c.rect(110, 34, 146, 84, "#FEF3C7", r=6)
        c.rect(120, 40, 136, 80, "#F59E0B", r=4)
        c.rect(106, 90, 150, 94, "#374151", r=1)
        c.rect(106, 99, 138, 102, "#9CA3AF", r=1)
        c.rect(106, 116, 150, 128, "#8135CD", r=6)
    else:
        for k in range(3):
            y = 34 + k * 16
            c.rect(106, y, 150, y + 3, "#9CA3AF", r=1)
        c.rect(106, 86, 150, 104, "#F5F3FF", r=4)
        c.rect(106, 114, 150, 128, "#16A34A" if i >= 50 else "#8135CD", r=6)
    targets = [(128, 45), (128, 45), (128, 122), (128, 122), (128, 121)]
    seg = min(4, i // 12)
    tx, ty = targets[seg]
    pulse = (i % 12) / 12
    r = 5 + pulse * 5
    c.ellipse(tx - r, ty - r, tx + r, ty + r, (129, 53, 205, int(160 * (1 - pulse))))
    c.ellipse(tx - 4, ty - 4, tx + 4, ty + 4, (129, 53, 205, 220))
    c.text((200, 118), "orgtik", 12, "#8135CD", bold=True, anchor="mm")
    return c.final()


def vid_demo(i):
    c = Canvas(VW, VH, bg="#FFFFFF")
    c.rect(0, 0, 256, 20, "#8135CD")
    c.text((8, 4), "Orgtik Analytics · Q3 2026", 9, "#FFFFFF", bold=True)
    values = [0.35, 0.5, 0.42, 0.66, 0.78, 0.9]
    t = min(1.0, i / 24)
    for k, v in enumerate(values):
        h = 90 * v * min(1.0, max(0.0, t * 1.3 - k * 0.06))
        x0 = 16 + k * 22
        c.rect(x0, 128 - h, x0 + 14, 128, "#8135CD" if k == 5 else "#C4B5FD", r=2)
    c.line([(12, 128), (150, 128)], "#E5E7EB", width=1)
    rev = int(21.4 * 10 * t) / 10
    c.text((168, 34), "Revenue", 9, "#6B7280")
    c.text((168, 46), f"EGP {rev:.1f}M", 15, "#111827", bold=True)
    c.text((168, 70), "Growth", 9, "#6B7280")
    c.text((168, 82), f"+{int(18 * t)}% YoY", 15, "#059669", bold=True)
    if i >= 36:
        c.rect(168, 108, 248, 126, "#F5F3FF", r=6)
        c.text((208, 117), "orgtik.test/demo", 8, "#8135CD", bold=True, anchor="mm")
    return c.final()


def build_videos(out_dir):
    for key, fn in (("mp4_ramadan_teaser", vid_ramadan), ("mp4_app_walkthrough", vid_walkthrough), ("mp4_product_demo", vid_demo)):
        frames = [fn(i) for i in range(60)]
        sps, pps, samples = h264.encode(frames, FPS, GOP)
        p = out_dir / f"{key}.mp4"
        p.write_bytes(h264.mp4(sps, pps, samples, VW, VH, FPS))
        playlist, segs = h264.hls(sps, pps, samples, FPS, GOP)
        hdir = ASSETS / "hls" / key
        hdir.mkdir(parents=True, exist_ok=True)
        (hdir / "index.m3u8").write_text(playlist, encoding="utf-8", newline="\n")
        assets = [{"key": "index.m3u8", "file": rel(hdir / "index.m3u8"), "mimeType": "application/vnd.apple.mpegurl", "size": (hdir / "index.m3u8").stat().st_size}]
        for n, s in enumerate(segs):
            sp = hdir / f"segment-{n}.ts"
            sp.write_bytes(s)
            assets.append({"key": sp.name, "file": rel(sp), "mimeType": "video/mp2t", "size": sp.stat().st_size})
        frames[40].save(hdir / "poster.png")
        register(key, p, "video/mp4", {"kind": "hls_video", "entryKey": "index.m3u8", "assets": assets})
        print("video", key, p.stat().st_size, [a["size"] for a in assets])


# ============================================================== main
def main():
    for sub in ("originals", "previews", "hls"):
        reset_dir(ASSETS / sub)
    out = ASSETS / "originals"
    for key, spec in C.PDFS.items():
        p = build_pdf(key, spec, out)
        register(key, p, "application/pdf", render_pages(key, p))
        print("pdf", key, p.stat().st_size, MANIFEST[key]["preview"]["pageCount"], "pages")
    for key, spec in C.DOCX.items():
        p, prev = build_docx(key, spec, out)
        register(key, p, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", prev)
        print("docx", key, p.stat().st_size)
    for key, sheets in C.XLSX.items():
        p, prev = build_xlsx(key, sheets, out)
        register(key, p, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", prev)
        print("xlsx", key, p.stat().st_size)
    for key, spec in C.PPTX.items():
        p, prev = build_pptx(key, spec, out)
        register(key, p, "application/vnd.openxmlformats-officedocument.presentationml.presentation", prev)
        print("pptx", key, p.stat().st_size)
    for key, lines in C.CSV.items():
        p = out / f"{key}.csv"
        p.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")
        rows = [ln.split(",") for ln in lines]
        register(key, p, "text/csv", preview_from_html(key, f'<div class="band"><b>ORGTIK</b> &#160; CSV preview</div>{table_html(rows)}', landscape=True))
    for key, text in C.TXT.items():
        p = out / f"{key}.txt"
        p.write_text(text, encoding="utf-8", newline="\n")
        body = "".join(f"<p>{esc(ln) or '&#160;'}</p>" for ln in text.splitlines())
        register(key, p, "text/plain", preview_from_html(key, f'<div style="font-size:10pt">{body}</div>'))
    images = {"png_insta_post": (img_insta_post, "PNG"), "png_linkedin_banner": (img_linkedin, "PNG"), "png_booth": (img_booth, "PNG"),
              "png_app_home": (img_app_home, "PNG"), "png_app_checkout": (img_app_checkout, "PNG"), "png_dashboard_dark": (img_dashboard_dark, "PNG"),
              "png_screenshot": (img_screenshot, "PNG"), "png_summer_banner": (img_summer_banner, "PNG"), "png_old_logo": (img_old_logo, "PNG"),
              "jpg_mango": (img_mango, "JPEG"), "jpg_family_pack": (img_family_pack, "JPEG"), "jpg_team_iftar": (img_iftar, "JPEG")}
    for key, (fn, fmt) in images.items():
        save_img(key, fn(), fmt, out)
        print("img", key, MANIFEST[key]["size"])
    build_misc(out)
    build_videos(out)
    (ASSETS / "manifest.json").write_text(json.dumps(MANIFEST, indent=1, ensure_ascii=False), encoding="utf-8")
    total = sum(f.stat().st_size for f in ASSETS.rglob("*") if f.is_file())
    print(f"\n{len(MANIFEST)} assets, {total / 1024 / 1024:.2f} MB total")


if __name__ == "__main__":
    main()
